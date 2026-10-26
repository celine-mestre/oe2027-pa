#!/usr/bin/env python3
"""
Sincroniza o ficheiro mestre da DACDP (Registo_PA_OE2027.xlsx, no SharePoint) com a base de dados das PA.

O que faz, na folha «Propostas de Alteração» (modelo da DACDP):
  - acrescenta uma linha por PA nova (ID sequencial);
  - atualiza, nas linhas existentes (chave: N.º PA), só as colunas que vêm da AR:
      C Data Entrada · E Partido · F Deputado/Proponente · G Artigo Alterado · S Estado · T Resultado Votação ·
      U Sentido de Voto · V Data Votação · L Descrição Sumária (só se estiver vazia) · W Observações (só acrescenta as ligações)
    e as colunas X em diante (informação adicional da UPE);
  - nunca altera as colunas de análise da DACDP (D, H a K, M a R, e o texto escrito em L e W);
  - preserva as outras folhas da DACDP («Medidas Anunciadas», …) e repõe as listas pendentes;
  - acrescenta à folha «Registo de Consultas» as verificações de fontes desde a última sincronização;
  - exporta para a base (data/oe<ano>/analise.csv) as classificações da DACDP que podem ir ao painel:
    área governativa, ministério, tipo de PA e tema. Relevância, impactos, análise jurídica e observações
    ficam só no ficheiro mestre (o repositório e o painel são públicos).

  python scripts/sincronizar_mestre.py "C:/…/Registo_PA_OE2027.xlsx"
  python scripts/sincronizar_mestre.py mestre.xlsx --ano 2027 --sem-analise      # não exporta analise.csv
  python scripts/sincronizar_mestre.py mestre.xlsx --ensaio                      # mostra o que faria, sem gravar

Guarda uma cópia de segurança do mestre (mestre.xlsx → mestre.YYYY-MM-DD_HHMM.bak.xlsx) antes de gravar.
"""
from __future__ import annotations

import argparse
import csv
from copy import copy
import datetime as dt
import json
import shutil
import sys
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "scripts"))
import exportar_excel as ex  # noqa: E402

FOLHA = "Propostas de Alteração"
COL = {n: i + 1 for i, n in enumerate(ex.COLUNAS_DACDP + ex.COLUNAS_UPE)}  # nome → índice (1-based)
# colunas da AR que a sincronização atualiza sempre (índice na linha gerada por exportar_excel.linha_pa)
ATUALIZAR = ["Data Entrada", "Partido", "Deputado/Proponente", "Artigo Alterado", "Estado", "Resultado Votação",
             "Sentido de Voto", "Data Votação"] + ex.COLUNAS_UPE
ANALISE = {"area_governativa": "Área Governativa", "ministerio": "Ministério", "tipo_pa": "Tipo de PA", "tema_dacdp": "Tema"}


def main(argv=None) -> int:
    from openpyxl import load_workbook
    from openpyxl.styles import Font

    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("mestre")
    ap.add_argument("--ano", default=ex.CONFIG["edicao_ativa"])
    ap.add_argument("--pasta", default="", help="outra pasta de dados (testes)")
    ap.add_argument("--sem-analise", action="store_true")
    ap.add_argument("--ensaio", action="store_true")
    a = ap.parse_args(argv)
    ed = ex.CONFIG["edicoes"][str(a.ano)]
    pasta = Path(a.pasta) if a.pasta else RAIZ / ed["pasta"]
    mestre = Path(a.mestre)
    if not mestre.exists():
        print(f"ficheiro mestre não encontrado: {mestre}", file=sys.stderr)
        return 2

    pa = ex.ler_json(pasta / "pa.json", [])
    classif = {c["n_pa"]: c for c in ex.ler_csv(pasta / "classificacao.csv")}
    wb = load_workbook(mestre)
    if FOLHA not in wb.sheetnames:
        print(f"a folha «{FOLHA}» não existe no mestre", file=sys.stderr)
        return 2
    ws = wb[FOLHA]
    cab = [str(c.value or "").strip() for c in ws[1]]
    if cab[:23] != ex.COLUNAS_DACDP:
        print("o cabeçalho das colunas A–W não corresponde ao modelo da DACDP:", cab[:23], file=sys.stderr)
        return 2
    # colunas UPE em falta (ficheiro v1 da DACDP só tem A–W): acrescentam-se à direita
    for i, nome in enumerate(ex.COLUNAS_UPE, 24):
        if len(cab) < i or cab[i - 1] != nome:
            c = ws.cell(row=1, column=i, value=nome)
            c.font, c.fill, c.alignment = copy(ws["A1"].font), copy(ws["A1"].fill), copy(ws["A1"].alignment)
            ws.column_dimensions[c.column_letter].width = ex.LARGURAS[i - 1]

    linhas = {}
    ultimo_id, ultima_linha = 0, 1
    for r in range(2, ws.max_row + 1):
        n = ws.cell(row=r, column=COL["N.º PA"]).value
        idv = ws.cell(row=r, column=COL["ID"]).value
        if n not in (None, ""):
            linhas[str(n).strip()] = r
            ultima_linha = r
        if isinstance(idv, (int, float)):
            ultimo_id = max(ultimo_id, int(idv))

    novas, atualizadas = [], 0
    for reg in sorted(pa, key=lambda r: (ex.re.match(r"\d*", r["n_pa"]).group(0).zfill(6), r["n_pa"])):
        valores = ex.linha_pa(0, reg, classif.get(reg["n_pa"], {}))
        n = reg["n_pa"]
        if n in linhas:
            r = linhas[n]
            mudou = False
            for nome in ATUALIZAR:
                novo = valores[COL[nome] - 1]
                c = ws.cell(row=r, column=COL[nome])
                if (c.value or None) != (novo or None) and not (c.value is None and novo in ("", None)):
                    c.value = novo
                    mudou = True
            if not ws.cell(row=r, column=COL["Descrição Sumária"]).value:
                ws.cell(row=r, column=COL["Descrição Sumária"], value=reg.get("objeto"))
            obs = ws.cell(row=r, column=COL["Observações"])
            liga = ex.observacoes(reg)
            if liga and (not obs.value or "AR: " not in str(obs.value)):
                obs.value = f"{obs.value} {liga}".strip() if obs.value else liga
            atualizadas += mudou
        else:
            ultima_linha += 1
            ultimo_id += 1
            valores[0] = ultimo_id
            for i, v in enumerate(valores, 1):
                ws.cell(row=ultima_linha, column=i, value=v)
            linhas[n] = ultima_linha
            novas.append(n)
    for r in range(2, ultima_linha + 1):
        for idx in (3, 22, 38, 39):
            ws.cell(row=r, column=idx).number_format = "DD/MM/YYYY"
        if novas and r > ultima_linha - len(novas):
            for c in ws[r]:
                c.font = Font(name=ex.FONTE, size=10)
    ws.data_validations.dataValidation = []
    if "Listas" in wb.sheetnames:
        wl = wb["Listas"]
        for i, (k, vals) in enumerate(ex.LISTAS.items(), 1):  # listas em falta no mestre (p. ex. «Área Governativa»)
            if wl.cell(row=1, column=i).value != k:
                wl.cell(row=1, column=i, value=k).font = Font(name=ex.FONTE, size=10, bold=True)
                for j, v in enumerate(vals, 2):
                    wl.cell(row=j, column=i, value=v).font = Font(name=ex.FONTE, size=10)
        ex.validacoes(ws)
    if ws.auto_filter.ref:
        ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(ex.COLUNAS_DACDP) + len(ex.COLUNAS_UPE)).column_letter}{max(ultima_linha, ex.LINHAS_MAX + 1)}"

    # registo de consultas
    n_cons = 0
    if "Registo de Consultas" in wb.sheetnames:
        wc = wb["Registo de Consultas"]
        existentes = {str(wc.cell(row=r, column=1).value or "") for r in range(2, wc.max_row + 1)}
        for linha in ex.consultas(pasta, ex.ler_json(pasta / "estado_fontes.json", {}), ex.ler_csv(pasta / "historico.csv"), ed)[::-1]:
            if linha[0] not in existentes:
                wc.append(linha)
                n_cons += 1
        wc.append([dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%d %H:%M UTC"), "Sincronização UPE do ficheiro mestre", "",
                   f"{len(novas)} PA novas acrescentadas; {atualizadas} linhas atualizadas a partir da AR.", len(novas)])

    # classificações da DACDP para o painel (só as que podem ser públicas)
    analise = []
    for n, r in linhas.items():
        linha = {"n_pa": n}
        for k, nome in ANALISE.items():
            linha[k] = str(ws.cell(row=r, column=COL[nome]).value or "").strip()
        if any(linha[k] for k in ANALISE):
            analise.append(linha)

    resumo = {"mestre": str(mestre), "pa_na_base": len(pa), "linhas_no_mestre": len(linhas), "novas": len(novas),
              "atualizadas": atualizadas, "consultas_acrescentadas": n_cons, "classificacoes_dacdp": len(analise)}
    if a.ensaio:
        print(json.dumps(resumo | {"ensaio": True, "novas_lista": novas[:20]}, ensure_ascii=False))
        return 0
    copia = mestre.with_name(f"{mestre.stem}.{dt.datetime.now().strftime('%Y-%m-%d_%H%M')}.bak{mestre.suffix}")
    shutil.copy2(mestre, copia)
    wb.save(mestre)
    if not a.sem_analise:
        with (pasta / "analise.csv").open("w", newline="", encoding="utf-8-sig") as fh:
            w = csv.DictWriter(fh, fieldnames=["n_pa"] + list(ANALISE), delimiter=";")
            w.writeheader()
            w.writerows(analise)
    print(json.dumps(resumo | {"copia_seguranca": str(copia)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
