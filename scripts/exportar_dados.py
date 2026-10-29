#!/usr/bin/env python3
"""
Exporta para Excel a base de dados das PA, só com a informação publicada pela Assembleia da República
(dados abertos, categoria «Orçamento do Estado»), tal como a ferramenta a recolhe e mostra no painel.

  python scripts/exportar_dados.py              # edição ativa → docs/PA_OE2027_dados_AR.xlsx
  python scripts/exportar_dados.py --ano 2026   # → docs/PA_OE2026_dados_AR.xlsx

Duas folhas:
  Propostas                uma linha por registo publicado pela AR (incluindo as partes 120C-1, 120C-2)
  Sobre este ficheiro      fonte, data dos dados, número de registos, ligações
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parents[1]
CONFIG = json.loads((RAIZ / "config.json").read_text(encoding="utf-8"))
AZUL, FONTE = "275673", "Arial"

# As mesmas colunas da exportação feita no painel («Exportar para Excel»)
COLUNAS = [
    ("n_pa", "N.º PA", 10), ("pa_base", "N.º base", 9), ("parte", "Parte", 7), ("id_ar", "ID AR", 9),
    ("data_entrada", "Data e hora de entrada", 17), ("proponente", "Proponente", 14), ("proponente_deputados", "Deputados", 30),
    ("apresentada", "Apresentada", 12), ("incide", "Incide", 11), ("artigo", "Artigo / mapa", 16),
    ("artigo_epigrafe", "Epígrafe do artigo", 35), ("diplomas", "Diplomas alterados", 35), ("tipo_ar", "Tipo", 22),
    ("objeto", "Objeto", 60), ("tema_ar", "Tema (AR)", 30), ("estado", "Estado", 16), ("estado_ar", "Estado (AR)", 26),
    ("data_votacao", "Data da votação", 13), ("sede_votacao", "Sede da votação", 13), ("resultado_votacao", "Resultado da votação", 26),
    ("votos_por_partido", "Sentido de voto por grupo parlamentar", 50), ("avocacao", "Avocação", 9), ("requerimentos", "Requerimentos", 35),
    ("parecer_ra", "Parecer Regiões Autónomas", 12), ("ministerio_ar", "Ministério (AR)", 16), ("n_votacoes", "N.º de votações", 9),
    ("url_detalhe", "Ligação AR", 30), ("url_pdf", "PDF", 30), ("primeira_captura", "Primeira captura", 17),
    ("ultima_captura", "Última captura", 17), ("fonte", "Fonte", 10), ("ausente_da_fonte", "Ausente da fonte", 17),
]
NUMERICOS = {"parte", "id_ar", "n_votacoes"}
DATAS = {"data_entrada", "primeira_captura", "ultima_captura", "ausente_da_fonte"}


def ordem(r: dict):
    import re
    m = re.match(r"(\d+)([A-Z]*)(?:-(\d+))?", r.get("n_pa") or "")
    return (int(m.group(1)) if m else 10**9, m.group(2) if m else "", int(m.group(3) or 0) if m else 0)


def valor(k: str, v):
    v = "" if v is None else v
    if k in NUMERICOS and str(v).isdigit():
        return int(v)
    if k in DATAS and v:
        return str(v).replace("T", " ").replace("Z", "")[:16]
    return v


def construir(ano: str, saida: Path, pasta: Path | None = None) -> dict:
    ed = CONFIG["edicoes"][ano]
    pasta = pasta or RAIZ / ed["pasta"]
    f = pasta / "pa.json"
    pa = sorted(json.loads(f.read_text(encoding="utf-8")) if f.exists() else [], key=ordem)
    execucao = json.loads((pasta / "ultima_execucao.json").read_text(encoding="utf-8")) if (pasta / "ultima_execucao.json").exists() else {}

    wb = Workbook()
    ws = wb.active
    ws.title = "Propostas"
    ws.append([c[1] for c in COLUNAS])
    for c in ws[1]:
        c.font = Font(name=FONTE, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for r in pa:
        ws.append([valor(k, r.get(k)) for k, _, _ in COLUNAS])
    for i, (_, _, w) in enumerate(COLUNAS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS))}{max(2, len(pa) + 1)}"
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=FONTE, size=10)

    info = wb.create_sheet("Sobre este ficheiro")
    momento = execucao.get("momento", "")
    linhas = [
        (f"Propostas de alteração ao Orçamento do Estado para {ed['ano_oe']} — {ed['designacao']}", ""),
        ("Fonte", "Assembleia da República, dados abertos (categoria «Orçamento do Estado», XVII Legislatura), "
                  f"ficheiro OEPropostasAlteracao{ed['ano_oe']}Or_json.txt"),
        ("Ligação", CONFIG["fontes"]["dados_abertos_pagina_oe"]),
        ("Dados de", momento.replace("T", " ").replace("Z", "")[:16] + " UTC" if momento else "—"),
        ("Ficheiro gerado em", dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%d %H:%M UTC")),
        ("Registos", len(pa)),
        ("Propostas distintas (sem partes)", len({r.get("pa_base") or r.get("n_pa") for r in pa})),
        ("Nota", "Cada registo numerado pela AR ocupa uma linha, incluindo as partes (120C-1, 120C-2) em que a AR divide uma proposta "
                 "que altera vários artigos ou alíneas. Todos os campos são publicados pela AR, exceto os de captura (primeira e última "
                 "captura, fonte, ausente da fonte). A coluna «Estado» agrupa o «Estado (AR)» publicado."),
        ("Painel", "https://upe-sggov.github.io/oe2027-pa/"),
        ("Tratamento", "Área de Suporte à Decisão, Secretaria-Geral do Governo"),
    ]
    for l in linhas:
        info.append(list(l))
    info["A1"].font = Font(name=FONTE, size=13, bold=True, color=AZUL)
    info.column_dimensions["A"].width, info.column_dimensions["B"].width = 30, 110
    for row in info.iter_rows(min_row=2):
        row[0].font = Font(name=FONTE, size=10, bold=True)
        row[1].font = Font(name=FONTE, size=10)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    saida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(saida)
    return {"ficheiro": str(saida), "registos": len(pa)}


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--ano", default=CONFIG["edicao_ativa"])
    ap.add_argument("--saida", default="")
    ap.add_argument("--pasta", default="", help="outra pasta de dados (testes)")
    a = ap.parse_args(argv)
    saida = Path(a.saida) if a.saida else RAIZ / "docs" / f"PA_OE{a.ano}_dados_AR.xlsx"
    if CONFIG["edicoes"][str(a.ano)].get("congelada") and saida.exists() and not a.saida:
        print(f"OE {a.ano}: edição congelada — mantém-se {saida.name}"); return 0
    print(json.dumps(construir(str(a.ano), saida, Path(a.pasta) if a.pasta else None), ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
